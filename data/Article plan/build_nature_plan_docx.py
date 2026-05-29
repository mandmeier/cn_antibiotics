#!/usr/bin/env python3
"""Build Nature_Sustainability_plan.docx from plan markdown files."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent.parent
PLAN = Path("/Users/kali/.cursor/plans/nature_sustainability_manuscript_4b10b05e.plan.md")
APPENDIX = ROOT / "sandbox" / "Nature_Sustainability_plan_appendix_yearbook.md"
OUT = ROOT / "sandbox" / "Nature_Sustainability_plan.docx"


def strip_frontmatter(text: str) -> str:
    if text.startswith("---"):
        end = text.find("---", 3)
        if end != -1:
            return text[end + 3 :].lstrip()
    return text


def add_md_to_doc(doc: Document, text: str) -> None:
    for line in text.splitlines():
        if not line.strip():
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("- "):
            p = doc.add_paragraph(line[2:].strip(), style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
        elif line.startswith("|") and "---" not in line:
            doc.add_paragraph(line.strip())
        elif line.startswith("```"):
            continue
        else:
            doc.add_paragraph(line)


def main() -> None:
    parts = []
    if PLAN.exists():
        parts.append(strip_frontmatter(PLAN.read_text(encoding="utf-8")))
    if APPENDIX.exists():
        parts.append("\n\n")
        parts.append(APPENDIX.read_text(encoding="utf-8"))
    audit = ROOT / "sandbox/linkage_pipeline/data/analysis_ready/yearbook_table_audit_summary.md"
    if audit.exists():
        parts.append("\n\n# Yearbook table audit (latest run)\n\n")
        parts.append(audit.read_text(encoding="utf-8"))

    doc = Document()
    doc.add_heading("Nature Sustainability Analysis — Manuscript Plan", 0)
    add_md_to_doc(doc, "\n".join(parts))
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
